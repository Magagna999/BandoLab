"""Crea db/import/modello-domande.xlsx (modello da compilare per caricare domande in blocco)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(__file__), '..', 'import', 'modello-domande.xlsx')
F = 'Arial'
head_font = Font(name=F, bold=True, color='FFFFFF', size=10)
head_fill = PatternFill('solid', fgColor='1F3A5F')
req_fill = PatternFill('solid', fgColor='FFF2CC')   # giallo: campo da compilare
ex_font = Font(name=F, italic=True, color='7F7F7F', size=10)
base = Font(name=F, size=10)
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ---------------- Istruzioni ----------------
ws = wb.active; ws.title = 'Istruzioni'
rows = [
 ('Modello per il caricamento delle domande in BandoLab', 'title'),
 ('', None),
 ('Come si usa', 'h'),
 ('1. Foglio «Tassonomia»: elenca le materie (colonne A-B) e gli argomenti (colonne D-F). Ogni argomento appartiene a una materia.', None),
 ('2. Foglio «Domande»: una riga per domanda. Le celle gialle sono obbligatorie. Le liste a tendina aiutano a non sbagliare.', None),
 ('3. La riga con «ESEMPIO» nella prima colonna è solo un modello di formato: viene ignorata, puoi cancellarla.', None),
 ('4. Salva il file e consegnalo: viene controllato (stesse regole del database) e poi caricato. Se ci sono errori ricevi l\'elenco con il numero di riga.', None),
 ('', None),
 ('Regole delle domande (sono vincoli del database)', 'h'),
 ('• Testo di almeno 10 caratteri. Quattro opzioni (A-D), tutte compilate. Una sola risposta corretta.', None),
 ('• La spiegazione è obbligatoria: serve all\'utente dopo la risposta.', None),
 ('• Fonte obbligatoria (norma, articolo, manuale). L\'indirizzo web è facoltativo.', None),
 ('• Difficoltà: 1 = facile, 2 = media, 3 = difficile (decisa da chi scrive, non calcolata).', None),
 ('• Ultima verifica: data in cui hai controllato la domanda sulla fonte (formato AAAA-MM-GG o data Excel). Prossima revisione facoltativa: se vuota, +12 mesi.', None),
 ('• ID domanda: facoltativo. Se vuoto viene assegnato automaticamente. Non riusare lo stesso ID in due righe.', None),
 ('• Due domande con lo stesso testo nello stesso file vengono segnalate come duplicate.', None),
 ('', None),
 ('Cosa succede dopo il caricamento', 'h'),
 ('Le domande arrivano come bozza (o in revisione) e diventano visibili agli utenti solo quando sono pubblicate da un revisore diverso dall\'autore.', None),
 ('Non copiare testi di quiz ufficiali o coperti da diritti: scrivi domande tue, con la fonte normativa.', None),
]
for i, (t, k) in enumerate(rows, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    c.font = Font(name=F, size=14, bold=True) if k == 'title' else Font(name=F, size=11, bold=True) if k == 'h' else base
ws.column_dimensions['A'].width = 120

# ---------------- Tassonomia ----------------
wt = wb.create_sheet('Tassonomia')
heads = {1: 'ID materia', 2: 'Nome materia', 4: 'ID argomento', 5: 'Nome argomento', 6: 'ID materia dell\'argomento'}
for col, h in heads.items():
    c = wt.cell(row=1, column=col, value=h); c.font = head_font; c.fill = head_fill; c.alignment = wrap; c.border = box
ex = {1: 'ESEMPIO', 2: 'Esempio: Diritto amministrativo', 4: 'ESEMPIO-ARG', 5: 'Esempio: Procedimento amministrativo', 6: 'ESEMPIO'}
for col, v in ex.items():
    c = wt.cell(row=2, column=col, value=v); c.font = ex_font; c.border = box
for r in range(3, 201):
    for col in (1, 2, 4, 5, 6):
        wt.cell(row=r, column=col).border = box; wt.cell(row=r, column=col).font = base
for col, w in {'A': 16, 'B': 38, 'C': 3, 'D': 18, 'E': 42, 'F': 24}.items():
    wt.column_dimensions[col].width = w
wt['H1'] = 'Gli ID sono sigle senza spazi, es. s-amm, t-proc. Una volta usati non vanno cambiati.'
wt['H1'].font = Font(name=F, italic=True, size=10)
wt.freeze_panes = 'A2'

# ---------------- Domande ----------------
wd = wb.create_sheet('Domande')
cols = [
 ('ID domanda (facoltativo)', 14, False), ('Testo della domanda', 60, True),
 ('Opzione A', 28, True), ('Opzione B', 28, True), ('Opzione C', 28, True), ('Opzione D', 28, True),
 ('Risposta corretta (A-D)', 12, True), ('Spiegazione', 50, True),
 ('ID materia', 16, True), ('ID argomento', 18, True), ('Difficoltà (1-3)', 11, True),
 ('Fonte', 28, True), ('Indirizzo web fonte (facoltativo)', 30, False),
 ('Ultima verifica', 14, True), ('Prossima revisione (facoltativa)', 16, False)]
for i, (h, w, req) in enumerate(cols, 1):
    c = wd.cell(row=1, column=i, value=h); c.font = head_font; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = box
    wd.column_dimensions[c.column_letter].width = w
wd.row_dimensions[1].height = 42
example = ['ESEMPIO', 'Se la legge non prevede un termine diverso, entro quanti giorni deve concludersi un procedimento amministrativo?',
           '30 giorni', '60 giorni', '90 giorni', '120 giorni', 'A',
           "L'art. 2 della L. 241/1990 fissa in 30 giorni il termine generale.", 'ESEMPIO', 'ESEMPIO-ARG', 1,
           'L. 241/1990, art. 2', 'https://www.normattiva.it/', '2026-10-27', '']
for i, v in enumerate(example, 1):
    c = wd.cell(row=2, column=i, value=v); c.font = ex_font; c.alignment = wrap; c.border = box
for r in range(3, 1003):
    for i, (h, w, req) in enumerate(cols, 1):
        c = wd.cell(row=r, column=i); c.font = base; c.alignment = wrap; c.border = box
        if req: c.fill = req_fill
wd.freeze_panes = 'C2'

def dv(formula, col, **kw):
    v = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                       errorTitle='Valore non valido', error=kw.get('err', 'Scegli un valore dalla lista.'))
    wd.add_data_validation(v); v.add(f'{col}3:{col}1002')
dv('"A,B,C,D"', 'G')
dv('"1,2,3"', 'K')
dv('=Tassonomia!$A$2:$A$200', 'I', err='Usa un ID materia presente nel foglio Tassonomia.')
dv('=Tassonomia!$D$2:$D$200', 'J', err='Usa un ID argomento presente nel foglio Tassonomia.')
wd['B1'].comment = Comment('Almeno 10 caratteri.', 'BandoLab')
wd['H1'].comment = Comment('Obbligatoria: viene mostrata all\'utente dopo la risposta.', 'BandoLab')
wb.move_sheet('Domande', offset=-1)  # ordine: Istruzioni, Domande, Tassonomia
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('creato', os.path.normpath(OUT))
