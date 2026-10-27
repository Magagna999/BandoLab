#!/usr/bin/env python3
"""Controlla un file Excel compilato (modello: db/import/modello-domande.xlsx) e genera l'SQL di caricamento.

Uso:
  python3 db/tools/excel-to-sql.py FILE.xlsx --author UUID [--env real|demo] [--status draft|review|published]
                                  [--reviewer UUID] [--start N] [--out FILE.sql]

Le regole replicano i vincoli del database. Se c'è anche un solo errore non viene scritto nulla.
"""
import argparse, datetime as dt, json, re, sys, uuid as uuidlib
from openpyxl import load_workbook

ap = argparse.ArgumentParser()
ap.add_argument('xlsx'); ap.add_argument('--author', required=True)
ap.add_argument('--env', default='real', choices=['real', 'demo'])
ap.add_argument('--status', default='draft', choices=['draft', 'review', 'published'])
ap.add_argument('--reviewer'); ap.add_argument('--start', type=int, default=1)
ap.add_argument('--out')
a = ap.parse_args()

def is_uuid(s):
    try: uuidlib.UUID(s); return True
    except Exception: return False
if not is_uuid(a.author): sys.exit('--author deve essere un UUID')
if a.status == 'published':
    if not a.reviewer or not is_uuid(a.reviewer): sys.exit('Per "published" serve --reviewer (UUID)')
    if a.reviewer == a.author: sys.exit('Il revisore deve essere diverso dall\'autore (vincolo del database)')

wb = load_workbook(a.xlsx, data_only=True)
for n in ('Domande', 'Tassonomia'):
    if n not in wb.sheetnames: sys.exit(f'Manca il foglio «{n}»')
errors, warnings = [], []
def err(sheet, row, msg): errors.append(f'{sheet} riga {row}: {msg}')

def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()
def date(v):
    if v is None or s(v) == '': return None
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    try: return dt.date.fromisoformat(s(v))
    except Exception: return 'bad'
def ident_ok(x): return bool(re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_\-]{0,39}', x))

# ---- Tassonomia ----
wt = wb['Tassonomia']
subjects, topics = {}, {}
for r in range(2, wt.max_row + 1):
    sid, sname = s(wt.cell(r, 1).value), s(wt.cell(r, 2).value)
    if sid == 'ESEMPIO' or sid == '' and sname == '': pass
    elif sid and sname:
        if not ident_ok(sid): err('Tassonomia', r, f'ID materia «{sid}» non valido (solo lettere, numeri, - e _)')
        elif sid in subjects: err('Tassonomia', r, f'materia «{sid}» duplicata')
        else: subjects[sid] = sname
    else: err('Tassonomia', r, 'materia incompleta (servono ID e nome)')
    tid, tname, tsub = s(wt.cell(r, 4).value), s(wt.cell(r, 5).value), s(wt.cell(r, 6).value)
    if tid == 'ESEMPIO-ARG' or tid == '' and tname == '' and tsub == '': continue
    if tid and tname and tsub:
        if not ident_ok(tid): err('Tassonomia', r, f'ID argomento «{tid}» non valido')
        elif tid in topics: err('Tassonomia', r, f'argomento «{tid}» duplicato')
        else: topics[tid] = (tname, tsub)
    else: err('Tassonomia', r, 'argomento incompleto (servono ID, nome e ID materia)')
for tid, (_, sub) in topics.items():
    if sub not in subjects: errors.append(f'Tassonomia: l\'argomento «{tid}» punta alla materia «{sub}», che non è elencata')

# ---- Domande ----
wd = wb['Domande']
questions, seen_ids, seen_text = [], set(), {}
counter = a.start
today = dt.date.today()
LET = {'A': 0, 'B': 1, 'C': 2, 'D': 3}
for r in range(2, wd.max_row + 1):
    vals = [wd.cell(r, c).value for c in range(1, 16)]
    if all(s(v) == '' for v in vals): continue
    if s(vals[0]).upper() == 'ESEMPIO': continue
    qid, text = s(vals[0]), s(vals[1])
    opts = [s(v) for v in vals[2:6]]
    ans, expl = s(vals[6]).upper(), s(vals[7])
    sub, top = s(vals[8]), s(vals[9])
    diff = s(vals[10]); src, url = s(vals[11]), s(vals[12])
    lv, nr = date(vals[13]), date(vals[14])
    n0 = len(errors)
    if len(text) < 10: err('Domande', r, 'il testo deve avere almeno 10 caratteri')
    for i, o in enumerate(opts):
        if not o: err('Domande', r, f'opzione {"ABCD"[i]} vuota')
    if len(set(o.lower() for o in opts if o)) < len([o for o in opts if o]): err('Domande', r, 'due opzioni sono identiche')
    if ans not in LET: err('Domande', r, 'la risposta corretta deve essere A, B, C o D')
    if not expl: err('Domande', r, 'la spiegazione è obbligatoria')
    if sub not in subjects: err('Domande', r, f'materia «{sub}» non presente nel foglio Tassonomia')
    if top not in topics: err('Domande', r, f'argomento «{top}» non presente nel foglio Tassonomia')
    elif sub in subjects and topics[top][1] != sub: err('Domande', r, f'l\'argomento «{top}» appartiene a «{topics[top][1]}», non a «{sub}»')
    if diff not in ('1', '2', '3'): err('Domande', r, 'la difficoltà deve essere 1, 2 o 3')
    if not src: err('Domande', r, 'la fonte è obbligatoria')
    if url and not re.match(r'https?://', url): err('Domande', r, 'l\'indirizzo web deve iniziare con http:// o https://')
    if lv is None: err('Domande', r, 'manca la data di ultima verifica')
    elif lv == 'bad': err('Domande', r, 'ultima verifica: data non valida (usa AAAA-MM-GG)')
    elif lv > today: err('Domande', r, 'ultima verifica: la data è nel futuro')
    if nr == 'bad': err('Domande', r, 'prossima revisione: data non valida')
    if qid:
        if not ident_ok(qid): err('Domande', r, f'ID «{qid}» non valido')
        elif qid in seen_ids: err('Domande', r, f'ID «{qid}» ripetuto')
        seen_ids.add(qid)
    key = re.sub(r'\W+', ' ', text.lower()).strip()
    if key in seen_text: err('Domande', r, f'testo duplicato della riga {seen_text[key]}')
    else: seen_text[key] = r
    if len(errors) == n0:
        if not nr: nr = lv.replace(year=lv.year + 1) if not (lv.month == 2 and lv.day == 29) else lv.replace(year=lv.year + 1, day=28)
        elif nr <= lv: err('Domande', r, 'la prossima revisione deve essere dopo l\'ultima verifica'); continue
        questions.append(dict(row=r, id=qid, text=text, options=opts, correct=LET[ans], explanation=expl, subject=sub, topic=top,
                              difficulty=int(diff), source=src, url=url or None, lv=lv.isoformat(), nr=nr.isoformat()))

if not questions and not errors: errors.append('Non ci sono domande da caricare')
# assegna ID automatici (r0001...), evitando quelli già usati nel file
used = {q['id'] for q in questions if q['id']}
for q in questions:
    if not q['id']:
        while f'r{counter:04d}' in used: counter += 1
        q['id'] = f'r{counter:04d}'; used.add(q['id']); counter += 1

print(f'Domande valide: {len(questions)} - materie: {len(subjects)} - argomenti: {len(topics)}')
if errors:
    print(f'\nERRORI ({len(errors)}): nessun file SQL scritto.')
    for e in errors[:200]: print(' -', e)
    sys.exit(1)

# ---- SQL ----
out = []
def load(table, rows):
    if not rows: return
    cols = list(dict.fromkeys(k for r in rows for k in r))
    L = ', '.join(cols)
    out.append(f"insert into {table} ({L}) select {L} from jsonb_populate_recordset(null::{table}, $j${json.dumps(rows, ensure_ascii=False)}$j$::jsonb) on conflict do nothing;")
env = a.env
out.append(f'-- Generato da db/tools/excel-to-sql.py da {a.xlsx} - ambiente {env}, stato {a.status}')
load('subjects', [dict(id=k, env=env, name=v) for k, v in subjects.items()])
load('topics', [dict(id=k, env=env, name=v[0], subject_id=v[1]) for k, v in topics.items()])
load('questions', [dict(id=q['id'], env=env) for q in questions])
now = dt.datetime.now(dt.timezone.utc).isoformat()
approved = a.status == 'published'
load('question_versions', [dict(env=env, question_id=q['id'], v=1, status=a.status, text=q['text'], options=q['options'], correct=q['correct'],
        explanation=q['explanation'], subject_id=q['subject'], topic_id=q['topic'], difficulty=q['difficulty'], source=q['source'],
        source_url=q['url'], last_verified=q['lv'], next_review=q['nr'], author_id=a.author,
        reviewer_id=a.reviewer if approved else None, change_note='Importato da Excel',
        approved_at=now if approved else None, published_at=now if approved else None) for q in questions])
if approved:
    ids = ','.join("'" + q['id'] + "'" for q in questions)
    out.append(f"update questions set current_version = 1 where env = '{env}' and id in ({ids});")
sql = '\n'.join(out) + '\n'
path = a.out or a.xlsx.rsplit('.', 1)[0] + '.sql'
open(path, 'w', encoding='utf-8').write(sql)
print(f'SQL scritto in {path} ({len(sql)//1024} KB)')
